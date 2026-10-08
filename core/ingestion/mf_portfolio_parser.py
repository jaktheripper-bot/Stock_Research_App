"""AMC Monthly Portfolio Disclosure Parser.

SEBI requires every AMC to publish month-end portfolio statements. AMCs publish them
as Excel/CSV workbooks in a broadly common layout: one sheet per scheme, a header row
(ISIN | Name of the Instrument | Industry / Rating | Quantity | Market Value | % to NAV),
and section rows ("Equity & Equity related", "Debt Instruments", "Money Market
Instruments", ...). This module parses that layout generically:

* Sniffs the real file format from bytes (AMCs ship xlsx content under .xls names).
* Locates columns from header text instead of fixed positions.
* Normalises % to NAV whether the AMC reports 2.93 or 0.0293.
* Classifies each line into the engine's holding types.
"""

import csv
import difflib
import io
import logging
import re
from datetime import date, datetime
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

logger = logging.getLogger(__name__)

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
_MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"], start=1)}

_STOP_MARKERS = ("notes:", "notes", "derivatives disclosure", "additional notes", "benchmark name")
_TOTAL_PREFIXES = ("total", "sub total", "subtotal", "grand total", "net assets", "grand")


# --------------------------------------------------------------------------- #
# Workbook loading
# --------------------------------------------------------------------------- #
def sniff_format(data: bytes) -> str:
    """Returns 'xlsx', 'xls', 'html' or 'csv' from file magic bytes (never trusts extension)."""
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    head = data[:512].lstrip().lower()
    if head.startswith(b"<!doctype html") or head.startswith(b"<html"):
        return "html"
    return "csv"


def load_sheets(data: bytes) -> Iterator[Tuple[str, Iterable[List[Any]]]]:
    """Yields (sheet_name, row_iterator) for every sheet in an xlsx/xls/csv payload."""
    fmt = sniff_format(data)
    if fmt == "html":
        raise ValueError("Expected a spreadsheet but received an HTML page (blocked or moved link).")

    if fmt == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for name in wb.sheetnames:
            ws = wb[name]
            yield name, (list(r) for r in ws.iter_rows(values_only=True))
        return

    if fmt == "xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        for sh in wb.sheets():
            yield sh.name, (sh.row_values(i) for i in range(sh.nrows))
        return

    text = data.decode("utf-8", errors="ignore")
    yield "csv", (row for row in csv.reader(io.StringIO(text)))


# --------------------------------------------------------------------------- #
# Cell helpers
# --------------------------------------------------------------------------- #
def _clean(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def _to_float(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _clean(v).replace(",", "").replace("%", "")
    if s in ("", "-", "NIL", "Nil", "nil", "NA", "N.A."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _strip_parenthetical(name: str) -> str:
    name = re.sub(r"\(.*", "", name)  # drop "(An open ended ...)" descriptions
    return name.strip(" -–:")


def normalize_scheme_name(name: str) -> str:
    """Canonical key for matching AMFI scheme names to workbook sheet titles."""
    n = _strip_parenthetical(_clean(name)).lower()
    n = re.split(r"\s+-\s+(direct|regular)", n)[0]
    n = n.replace("&", " and ").replace("’", "'")
    n = re.sub(r"\b(the|mutual|fund)\b", " ", n)
    return re.sub(r"[^a-z0-9]+", "", n)


def match_scheme_name(target: str, candidates: List[str], cutoff: float = 0.93) -> Optional[str]:
    """Returns the single best candidate for target, or None (never guesses on ties)."""
    key = normalize_scheme_name(target)
    if not key:
        return None
    exact = [c for c in candidates if normalize_scheme_name(c) == key]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return None
    scored = sorted(
        ((difflib.SequenceMatcher(None, key, normalize_scheme_name(c)).ratio(), c) for c in candidates),
        reverse=True,
    )
    if not scored or scored[0][0] < cutoff:
        return None
    if len(scored) > 1 and (scored[0][0] - scored[1][0]) < 0.02:
        return None
    return scored[0][1]


# --------------------------------------------------------------------------- #
# Column + section detection
# --------------------------------------------------------------------------- #
def _detect_columns(row: List[Any]) -> Optional[Dict[str, int]]:
    cols: Dict[str, int] = {}
    for i, cell in enumerate(row):
        t = _clean(cell).lower()
        if not t:
            continue
        if t == "isin" or t.startswith("isin"):
            cols.setdefault("isin", i)
        elif "name of the instrument" in t or t in ("instrument", "security name", "name of instrument", "issuer"):
            cols.setdefault("name", i)
        elif "industry" in t or "rating" in t:
            cols.setdefault("sector", i)
        elif t.startswith("quantity") or t.startswith("qty"):
            cols.setdefault("qty", i)
        elif "market" in t and "value" in t or t.startswith("market/fair"):
            cols.setdefault("value", i)
        elif ("% to" in t and ("nav" in t or "net" in t)) or "% of net" in t or "% to net" in t or "%age to nav" in t:
            cols.setdefault("weight", i)
        elif t.startswith("yield") or t.startswith("ytm"):
            cols.setdefault("yield", i)
    if "isin" in cols and "weight" in cols:
        cols.setdefault("name", cols["isin"] + 1)
        return cols
    return None


def classify_section(text: str) -> Optional[str]:
    """Maps a section-title row to a holding type. None means 'not a recognised section'."""
    t = text.lower()
    if any(k in t for k in ("triparty", "tri-party", "treps", "reverse repo", "repo", "cash", "net current", "net receivable")):
        return "CASH_EQUIVALENT"
    if any(k in t for k in ("foreign", "overseas", "adr", "gdr")):
        return "FOREIGN_EQUITY"
    if "equity" in t and "related" in t:
        return "EQUITY"
    if t.startswith("equity") or "listed / awaiting listing on stock" in t and "debt" not in t and "equity" in t:
        return "EQUITY"
    if any(k in t for k in ("mutual fund unit", "exchange traded fund", "etf", "reit", "invit", "units of")):
        return "OTHER"
    if any(k in t for k in ("future", "option", "derivative")):
        return "DERIVATIVES"
    if any(k in t for k in (
        "government securit", "g-sec", "treasury bill", "t-bill", "state development", "debt",
        "debenture", "bond", "certificate of deposit", "commercial paper", "money market",
        "securitised", "securitized", "zero coupon", "pass through", "ncd",
    )):
        return "DEBT"
    return None


def _holding_type_from_isin(isin: str, section: Optional[str]) -> str:
    if section in ("EQUITY", "FOREIGN_EQUITY", "DEBT", "OTHER", "DERIVATIVES"):
        base = section
    elif isin.startswith("INF"):
        base = "OTHER"
    elif isin.startswith("INE") and isin[8:10] == "01":
        base = "EQUITY"
    elif isin.startswith("IN"):
        base = "DEBT"
    else:
        base = "FOREIGN_EQUITY"
    # An overseas ISIN inside an equity section is a foreign equity holding.
    if base == "EQUITY" and not isin.startswith("IN"):
        return "FOREIGN_EQUITY"
    return base


def parse_as_on_date(text: str) -> Optional[str]:
    """Parses 'as on August 31,2026' / '31-Aug-2026' / '31/08/2026' into ISO date."""
    t = _clean(text)
    m = re.search(r"([A-Za-z]{3,9})\.?\s+(\d{1,2})\s*,?\s*(\d{4})", t)
    if m and m.group(1).lower() in _MONTHS:
        return date(int(m.group(3)), _MONTHS[m.group(1).lower()], int(m.group(2))).isoformat()
    m = re.search(r"(\d{1,2})[-\s/]([A-Za-z]{3,9})[-\s/,]*(\d{2,4})", t)
    if m:
        mon = next((v for k, v in _MONTHS.items() if k.startswith(m.group(2).lower()[:3])), None)
        if mon:
            yr = int(m.group(3))
            yr += 2000 if yr < 100 else 0
            return date(yr, mon, int(m.group(1))).isoformat()
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", t)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return None
    return None


# --------------------------------------------------------------------------- #
# Sheet parsing
# --------------------------------------------------------------------------- #
def parse_scheme_sheet(rows: Iterable[List[Any]]) -> Optional[Dict[str, Any]]:
    """Parses one scheme sheet. Returns None if the sheet is not a portfolio statement."""
    scheme_title = ""
    as_on: Optional[str] = None
    cols: Optional[Dict[str, int]] = None
    section: Optional[str] = None
    holdings: List[Dict[str, Any]] = []
    header_seen_at = None

    for idx, row in enumerate(rows):
        texts = [_clean(c) for c in row]
        nonempty = [t for t in texts if t]
        if not nonempty:
            continue

        if cols is None:
            joined = " ".join(nonempty)
            if "portfolio statement" in joined.lower() or "as on" in joined.lower():
                as_on = as_on or parse_as_on_date(joined)
            if idx < 4 and not scheme_title and "portfolio statement" not in joined.lower():
                longest = max(nonempty, key=len)
                if len(longest) > 8 and not ISIN_RE.match(longest):
                    scheme_title = longest
            detected = _detect_columns(list(row))
            if detected:
                cols = detected
                header_seen_at = idx
            continue

        if len(nonempty) == 1:
            title = nonempty[0]
            low = title.lower()
            if any(low.startswith(s) for s in _STOP_MARKERS):
                break
            sec = classify_section(title)
            if sec:
                section = sec
            continue

        isin_cell = _clean(row[cols["isin"]]) if cols["isin"] < len(row) else ""
        weight = _to_float(row[cols["weight"]]) if cols["weight"] < len(row) else None
        name_cell = _clean(row[cols["name"]]) if cols["name"] < len(row) else ""

        if ISIN_RE.match(isin_cell.upper()) and weight is not None:
            isin = isin_cell.upper()
            sector = _clean(row[cols["sector"]]) if "sector" in cols and cols["sector"] < len(row) else ""
            htype = _holding_type_from_isin(isin, section)
            details: Dict[str, Any] = {}
            if "yield" in cols and cols["yield"] < len(row):
                y = _to_float(row[cols["yield"]])
                if y is not None:
                    details["ytm_raw"] = y
            if "qty" in cols and cols["qty"] < len(row):
                q = _to_float(row[cols["qty"]])
                if q is not None:
                    details["quantity"] = q
            if "value" in cols and cols["value"] < len(row):
                v = _to_float(row[cols["value"]])
                if v is not None:
                    details["market_value_lakhs"] = v
            holdings.append({
                "isin": isin,
                "holding_name": re.sub(r"\*+$", "", name_cell).strip(),
                "holding_type": htype,
                "sector_or_rating": sector,
                "weight_raw": weight,
                "instrument_details": details,
            })
            continue

        # Non-ISIN lines: only TREPS / repo / net current assets style cash rows are kept.
        label = name_cell or next((t for t in texts if t and _to_float(t) is None), "")
        if not label or any(label.lower().startswith(p) for p in _TOTAL_PREFIXES):
            continue
        if weight is None:
            continue
        if section == "CASH_EQUIVALENT" or classify_section(label) == "CASH_EQUIVALENT":
            holdings.append({
                "isin": "",
                "holding_name": label,
                "holding_type": "CASH_EQUIVALENT",
                "sector_or_rating": "CASH",
                "weight_raw": weight,
                "instrument_details": {},
            })

    if cols is None or not holdings:
        return None

    return {
        "scheme_title": scheme_title,
        "scheme_name": _strip_parenthetical(scheme_title),
        "as_on_date": as_on,
        "holdings": holdings,
        "header_row": header_seen_at,
    }


def normalize_weights(holdings: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Converts weight_raw into weight_pct. Handles fraction (0.0293) vs percent (2.93) files.

    Short futures in arbitrage funds are negative and are dropped: the engine treats weights as
    non-negative exposures, and the hedged equity leg is already captured as a long position.
    """
    positives = [h for h in holdings if h["weight_raw"] > 0]
    total = sum(h["weight_raw"] for h in positives)
    scale = 100.0 if 0.3 <= total <= 1.6 else 1.0
    out = []
    for h in positives:
        item = dict(h)
        item["weight_pct"] = round(h["weight_raw"] * scale, 4)
        if "ytm_raw" in item["instrument_details"]:
            y = item["instrument_details"].pop("ytm_raw")
            item["instrument_details"]["ytm"] = round(y * 100.0 if y < 1.0 else y, 4)
        out.append(item)
    return out


def parse_workbook(data: bytes) -> Dict[str, Dict[str, Any]]:
    """Parses every scheme sheet in a workbook. Returns {sheet_name: parsed_sheet}."""
    results: Dict[str, Dict[str, Any]] = {}
    for name, rows in load_sheets(data):
        try:
            parsed = parse_scheme_sheet(rows)
        except Exception as e:
            logger.warning(f"Skipping unparseable sheet '{name}': {e}")
            continue
        if parsed:
            parsed["holdings"] = normalize_weights(parsed["holdings"])
            if parsed["holdings"]:
                results[name] = parsed
    return results
