import io
import unittest
from datetime import date
from unittest.mock import patch

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    openpyxl = None

from core.ingestion.mf_portfolio_ingest import (
    amc_key_for_fund_house,
    build_holdings,
    ingest_top_funds,
    load_sources,
    recent_month_ends,
    render_url,
    validate_holdings,
)
from core.ingestion.mf_portfolio_parser import (
    match_scheme_name,
    normalize_scheme_name,
    parse_as_on_date,
    parse_workbook,
    sniff_format,
)
from core.ingestion.mf_universe import rank_funds_by_aum

HEADER = ["ISIN", "Name of the Instrument", "Industry / Rating", "Quantity",
          "Market/Fair Value\n( Rs. in Lacs)", "% to NAV"]


def _equity_sheet(wb, title, scheme_title, weights, as_fraction=True):
    ws = wb.create_sheet(title)
    ws.append(["RLMF001", scheme_title, "Index"])
    ws.append(["Monthly Portfolio Statement as on August 31,2026"])
    ws.append([])
    ws.append([None] + HEADER[:])  # leading code column shifts headers by one, like real files
    ws.append(["Equity & Equity related"])
    ws.append(["(a) Listed / awaiting listing on Stock Exchanges"])
    isins = ["INE040A01034", "INE002A01018", "INE009A01021"]
    for i, w in enumerate(weights):
        ws.append(["CODE%d" % i, isins[i], "Company %d Limited**" % i, "Banks", 1000, 500.0,
                   w if as_fraction else w * 100])
    ws.append(["Total", None, None, None, None, 3000.0, sum(weights)])
    ws.append(["Money Market Instruments"])
    ws.append(["Triparty Repo/ Reverse Repo Instrument"])
    ws.append(["Triparty Repo", None, None, None, None, 100.0, round(1 - sum(weights) - 0.001, 4) if as_fraction
               else round((1 - sum(weights) - 0.001) * 100, 2)])
    ws.append(["Total", None, None, None, None, 100.0, 0.02])
    ws.append(["OTHERS"])
    ws.append(["Net Current Assets", None, None, None, None, 10.0, 0.001 if as_fraction else 0.1])
    ws.append(["Notes:"])
    ws.append(["Note one with an INE000000000 mention"])
    return ws


def _book_bytes(builder):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    builder(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is required for synthetic workbook ingestion tests")
class TestParser(unittest.TestCase):
    def test_sniff_ignores_extension(self):
        data = _book_bytes(lambda wb: _equity_sheet(wb, "EA", "Alpha Fund", [0.5, 0.3, 0.17]))
        self.assertEqual(sniff_format(data), "xlsx")
        self.assertEqual(sniff_format(b"<html><body>blocked</body></html>"), "html")
        self.assertEqual(sniff_format(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest"), "xls")

    def test_parses_fraction_weights_and_cash(self):
        data = _book_bytes(lambda wb: _equity_sheet(
            wb, "EA", "Alpha Equity Fund (An open ended equity scheme)", [0.5, 0.3, 0.17]))
        parsed = parse_workbook(data)["EA"]
        self.assertEqual(parsed["scheme_name"], "Alpha Equity Fund")
        self.assertEqual(parsed["as_on_date"], "2026-08-31")
        hs = parsed["holdings"]
        self.assertAlmostEqual(sum(h["weight_pct"] for h in hs), 100.0, delta=0.01)
        self.assertEqual(sum(1 for h in hs if h["holding_type"] == "EQUITY"), 3)
        cash = [h for h in hs if h["holding_type"] == "CASH_EQUIVALENT"]
        self.assertEqual({c["holding_name"] for c in cash}, {"Triparty Repo", "Net Current Assets"})
        self.assertTrue(all(h["holding_name"].endswith("Limited") for h in hs if h["holding_type"] == "EQUITY"))

    def test_parses_percent_weights(self):
        data = _book_bytes(lambda wb: _equity_sheet(wb, "EB", "Beta Fund", [0.5, 0.3, 0.17], as_fraction=False))
        hs = parse_workbook(data)["EB"]["holdings"]
        self.assertAlmostEqual(sum(h["weight_pct"] for h in hs), 100.0, delta=0.05)

    def test_note_text_not_parsed_as_holdings(self):
        data = _book_bytes(lambda wb: _equity_sheet(wb, "EA", "Alpha Fund", [0.5, 0.3, 0.17]))
        hs = parse_workbook(data)["EA"]["holdings"]
        self.assertFalse(any("INE000000000" in h["isin"] for h in hs))

    def test_non_portfolio_sheets_ignored(self):
        def build(wb):
            ws = wb.create_sheet("Index")
            ws.append(["Code", "Scheme"])
            _equity_sheet(wb, "EA", "Alpha Fund", [0.5, 0.3, 0.17])
        self.assertEqual(list(parse_workbook(_book_bytes(build)).keys()), ["EA"])

    def test_html_rejected(self):
        with self.assertRaises(ValueError):
            parse_workbook(b"<!DOCTYPE html><html></html>")

    def test_date_formats(self):
        self.assertEqual(parse_as_on_date("Monthly Portfolio Statement as on August 31,2026"), "2026-08-31")
        self.assertEqual(parse_as_on_date("as on 31-Aug-2026"), "2026-08-31")
        self.assertEqual(parse_as_on_date("Portfolio as on 31/08/2026"), "2026-08-31")
        self.assertIsNone(parse_as_on_date("no date here"))

    def test_scheme_name_matching(self):
        titles = ["Nippon India Small Cap Fund", "Nippon India Large Cap Fund", "Nippon India Multi Cap Fund"]
        self.assertEqual(
            match_scheme_name("Nippon India Small Cap Fund - Direct Plan Growth Plan - Growth Option", titles),
            "Nippon India Small Cap Fund")
        self.assertIsNone(match_scheme_name("Nippon India Gold Savings Fund - Direct Plan - Growth", titles))
        self.assertEqual(normalize_scheme_name("Alpha & Beta  Fund (An open ended)"), "alphaandbeta")

    def test_ambiguous_match_returns_none(self):
        self.assertIsNone(match_scheme_name("Alpha Fund", ["Alpha Fund", "ALPHA FUND"]))


class TestIngestionHelpers(unittest.TestCase):
    def test_month_end_and_url(self):
        ends = recent_month_ends(3, date(2026, 10, 8))
        self.assertEqual(ends, [date(2026, 9, 30), date(2026, 8, 31), date(2026, 7, 31)])
        url = render_url(load_sources()["nippon"]["url_template"], date(2026, 8, 31))
        self.assertTrue(url.endswith("NIMF-MONTHLY-PORTFOLIO-31-Aug-26.xls"))

    def test_registry_only_enables_verified_amcs(self):
        for key, cfg in load_sources().items():
            if cfg.get("url_template"):
                self.assertTrue(cfg.get("verified_on"), f"{key} template enabled without verification date")

    def test_fund_house_to_key(self):
        src = load_sources()
        self.assertEqual(amc_key_for_fund_house("Nippon India Mutual Fund", src), "nippon")
        self.assertIsNone(amc_key_for_fund_house("Unknown MF", src))

    def test_validate_rejects_incomplete(self):
        self.assertIsNotNone(validate_holdings([]))
        self.assertIsNotNone(validate_holdings([{"weight_pct": 60.0}]))
        self.assertIsNone(validate_holdings([{"weight_pct": 99.2}]))

    def test_build_holdings_maps_isin_and_merges(self):
        parsed = {
            "as_on_date": "2026-08-31",
            "holdings": [
                {"isin": "INE040A01034", "holding_name": "HDFC Bank", "holding_type": "EQUITY",
                 "sector_or_rating": "Banks", "weight_pct": 5.0, "instrument_details": {}},
                {"isin": "INE040A01034", "holding_name": "HDFC Bank", "holding_type": "EQUITY",
                 "sector_or_rating": "Banks", "weight_pct": 1.0, "instrument_details": {}},
                {"isin": "", "holding_name": "Triparty Repo", "holding_type": "CASH_EQUIVALENT",
                 "sector_or_rating": "CASH", "weight_pct": 2.0, "instrument_details": {}},
                {"isin": "IN0020160035", "holding_name": "6.97% GOI", "holding_type": "DEBT",
                 "sector_or_rating": "SOVEREIGN", "weight_pct": 3.0, "instrument_details": {"ytm": 5.2}},
            ],
        }
        out = build_holdings(parsed, {"INE040A01034": "HDFCBANK"})
        by_id = {h["identifier"]: h for h in out}
        self.assertEqual(by_id["HDFCBANK"]["weight_pct"], 6.0)
        self.assertIn("TREPS", by_id)
        self.assertIn("IN0020160035", by_id)
        self.assertEqual(by_id["IN0020160035"]["instrument_details"]["ytm"], 5.2)
        self.assertEqual(out[0]["identifier"], "HDFCBANK")


class TestUniverseRanking(unittest.TestCase):
    AAUM = [
        {"Mfname": "Alpha Mutual Fund", "SchemeCat_Desc": "Equity Scheme - Flexi Cap Fund", "schemes": [
            {"SchemeNAVName": "Alpha Flexi Fund - Direct Plan - Growth", "AMFI_Code": 1,
             "AverageAumForTheMonth": {"ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas": 5000000.0}},
            {"SchemeNAVName": "Alpha Flexi Fund - Regular Plan - Growth", "AMFI_Code": 2,
             "AverageAumForTheMonth": {"ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas": 9000000.0}},
            {"SchemeNAVName": "Alpha Flexi Fund - Direct Plan - IDCW Payout", "AMFI_Code": 3,
             "AverageAumForTheMonth": {"ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas": 9000000.0}},
        ]},
        {"Mfname": "Beta Mutual Fund", "SchemeCat_Desc": "Debt Scheme - Liquid Fund", "schemes": [
            {"SchemeNAVName": "Beta Liquid Fund - Direct Plan - Growth", "AMFI_Code": 4,
             "AverageAumForTheMonth": {"ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas": 7000000.0}},
        ]},
        {"Mfname": "Beta Mutual Fund", "SchemeCat_Desc": "Other Scheme - FoF Domestic", "schemes": [
            {"SchemeNAVName": "Beta FoF - Direct Plan - Growth", "AMFI_Code": 5,
             "AverageAumForTheMonth": {"ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas": 99999999.0}},
        ]},
    ]

    def test_ranks_direct_growth_only_and_converts_to_crores(self):
        ranked = rank_funds_by_aum(self.AAUM)
        self.assertEqual([r["scheme_code"] for r in ranked], ["4", "1"])
        self.assertEqual(ranked[0]["aum_crores"], 70000.0)
        self.assertEqual([r["aum_rank"] for r in ranked], [1, 2])


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is required for synthetic workbook ingestion tests")
class TestOrchestratorSkipAndReplace(unittest.TestCase):
    def _universe(self):
        return [
            {"aum_rank": 1, "scheme_code": "9001", "scheme_name": "Nippon India Alpha Fund - Direct Plan - Growth",
             "fund_house": "Nippon India Mutual Fund", "category": "Flexi Cap Fund",
             "broad_category": "EQUITY", "aum_crores": 900.0},
            {"aum_rank": 2, "scheme_code": "9002", "scheme_name": "SBI Beta Fund - Direct Plan - Growth",
             "fund_house": "SBI Mutual Fund", "category": "Flexi Cap Fund",
             "broad_category": "EQUITY", "aum_crores": 800.0},
            {"aum_rank": 3, "scheme_code": "9003", "scheme_name": "Nippon India Missing Fund - Direct Plan - Growth",
             "fund_house": "Nippon India Mutual Fund", "category": "Flexi Cap Fund",
             "broad_category": "EQUITY", "aum_crores": 700.0},
            {"aum_rank": 4, "scheme_code": "9004", "scheme_name": "Nippon India Gamma Fund - Direct Plan - Growth",
             "fund_house": "Nippon India Mutual Fund", "category": "Flexi Cap Fund",
             "broad_category": "EQUITY", "aum_crores": 600.0},
        ]

    def _provider(self):
        def build(wb):
            _equity_sheet(wb, "A1", "Nippon India Alpha Fund (An open ended scheme)", [0.5, 0.3, 0.17])
            _equity_sheet(wb, "A2", "Nippon India Gamma Fund (An open ended scheme)", [0.6, 0.2, 0.19])
        data = _book_bytes(build)
        return lambda key, cfg: (data, "test") if key == "nippon" else None

    @patch("core.ingestion.mf_portfolio_ingest.fetch_amfi_nav_raw", return_value="")
    def test_skips_funds_without_data_and_replaces_with_next(self, _):
        report = ingest_top_funds(
            target=2, dry_run=True, universe=self._universe(),
            workbook_provider=self._provider(), isin_map={})
        self.assertEqual([f["scheme_code"] for f in report["ingested"]], ["9001", "9004"])
        self.assertTrue(report["target_met"])
        reasons = {s["scheme_code"]: s["reason"] for s in report["skipped"]}
        self.assertIn("no workbook", reasons["9002"])
        self.assertIn("no matching scheme sheet", reasons["9003"])

    @patch("core.ingestion.mf_portfolio_ingest.fetch_amfi_nav_raw", return_value="")
    def test_reports_shortfall_honestly(self, _):
        report = ingest_top_funds(
            target=10, dry_run=True, universe=self._universe(),
            workbook_provider=self._provider(), isin_map={})
        self.assertFalse(report["target_met"])
        self.assertEqual(report["ingested_count"], 2)

    @patch("core.ingestion.mf_portfolio_ingest.fetch_amfi_nav_raw", return_value="")
    def test_dry_run_writes_nothing(self, _):
        with patch("core.ingestion.mf_portfolio_ingest.save_mutual_fund_holdings") as save_h, \
             patch("core.ingestion.mf_portfolio_ingest.save_mutual_fund_scheme") as save_s:
            ingest_top_funds(target=2, dry_run=True, universe=self._universe(),
                             workbook_provider=self._provider(), isin_map={})
            save_h.assert_not_called()
            save_s.assert_not_called()

    @patch("core.ingestion.mf_portfolio_ingest.fetch_amfi_nav_raw", return_value="")
    def test_real_run_saves_holdings_and_scheme(self, _):
        with patch("core.ingestion.mf_portfolio_ingest.save_mutual_fund_holdings", return_value=True) as save_h, \
             patch("core.ingestion.mf_portfolio_ingest.save_mutual_fund_scheme", return_value=True) as save_s, \
             patch("core.ingestion.mf_portfolio_ingest.get_mutual_fund_scheme", return_value=None):
            ingest_top_funds(target=1, dry_run=False, universe=self._universe(),
                             workbook_provider=self._provider(), isin_map={"INE040A01034": "HDFCBANK"})
            self.assertEqual(save_h.call_count, 1)
            code, holdings = save_h.call_args[0]
            self.assertEqual(code, "9001")
            self.assertIn("HDFCBANK", [h["identifier"] for h in holdings])
            saved_scheme = save_s.call_args[0][0]
            self.assertEqual(saved_scheme["aum_crores"], 900.0)
            self.assertEqual(saved_scheme["metadata"]["holdings_as_on"], "2026-08-31")


if __name__ == "__main__":
    unittest.main()
